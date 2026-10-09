# Referencia: cálculo independiente de los resultados esperados del servicio (Python + openpyxl).
import openpyxl, json, datetime as dt, unicodedata, re, sys
def norm(s): return unicodedata.normalize('NFD',str(s or '')).encode('ascii','ignore').decode().lower().strip()
def r2(x): return round(x+1e-9,2)
ALIAS={'subtotal':['sub total','subtotal'],'total':['total'],'iva':['total iva tras.'],'descuento':['descuento'],'estatus':['estatus'],
 'fecha':['fecha'],'folio':['folio'],'serie':['serie'],'rfc':['r.f.c. receptor'],'receptor':['razon social receptor'],'tipo':['tipo de comprobante'],
 'ivaret':['total iva ret.'],'isrret':['total isr ret.']}
def load(path):
    ws=openpyxl.load_workbook(path,data_only=True).worksheets[0]
    h={}
    for c in ws[1]:
        k=norm(c.value)
        for key,names in ALIAS.items():
            if key not in h and k in names: h[key]=c.column
    last=max(c.column for c in ws[1] if c.value not in (None,''))
    rows=[]
    for r in range(2,ws.max_row+1):
        f=ws.cell(r,h['fecha']).value; fo=ws.cell(r,h['folio']).value
        if f is None and fo is None: continue
        g=lambda k:(ws.cell(r,h[k]).value if k in h else None)
        n=lambda k: float(g(k) or 0)
        x=dict(r=r,folio=str(fo),serie=str(g('serie') or ''),dia=f.date().isoformat(),sub=n('subtotal'),tot=n('total'),iva=n('iva'),desc=n('descuento'),
               est=str(g('estatus') or ''),tipo=str(g('tipo') or 'INGRESO').upper(),rfc=str(g('rfc') or ''),ivaret=n('ivaret'),isrret=n('isrret'))
        e=r2(x['tot']-(x['sub']-x['desc'])-x['iva']+x['ivaret']+x['isrret']); x['ieps']=r2(e) if e>0.02 else 0
        rows.append(x)
    return h,last,rows
def golden(path,marcados_folios=(),tol=1.0):
    h,last,rows=load(path)
    act=lambda x: norm(x['est'])=='activo' and x['tipo'] in ('INGRESO','I','')
    dias=[]; i=0
    while i<len(rows):
        j=i
        while j+1<len(rows) and rows[j+1]['dia']==rows[i]['dia']: j+=1
        blk=rows[i:j+1]; a=[x for x in blk if act(x)]
        d=dict(dia=rows[i]['dia'],fila_ini=rows[i]['r'],fila_fin=rows[j]['r'],n=len(blk),cancelados=sum(1 for x in blk if norm(x['est'])=='cancelado'),
               S=r2(sum(x['sub'] for x in a)),T=r2(sum(x['tot'] for x in a)),U=r2(sum(x['iva'] for x in a)),V=r2(sum(x['desc'] for x in a)),IEPS=r2(sum(x['ieps'] for x in a)))
        d['W_exacta']=(d['S']-d['V'])*0.16-d['U']; d['W']=r2(d['W_exacta']); dias.append(d); i=j+1
    cand=[]
    for x in rows:
        if act(x) and (x['sub']-x['desc']+x['ieps'])*0.16-x['iva']>tol:
            cand.append(dict(fila=x['r'],folio=x['folio'],dia=x['dia'],subtotal=x['sub'],iva=x['iva'],ieps=x['ieps'],base0=r2(x['sub']-x['desc']+x['ieps']-x['iva']/0.16),mixto=x['iva']>0))
    marc={c['fila'] for c in cand if c['folio'] in marcados_folios}
    for d in dias:
        X=sum(c['base0'] for c in cand if c['fila'] in marc and d['fila_ini']<=c['fila']<=d['fila_fin'])
        d['X']=r2(X); d['revision']='OK' if abs(d['W_exacta']-(X-d['IEPS'])*0.16)<=tol else 'REVISAR'
    ieps=[dict(fila=x['r'],folio=x['folio'],dia=x['dia'],ieps=x['ieps']) for x in rows if act(x) and x['ieps']>0]
    canc=[]
    for x in rows:
        if norm(x['est'])=='cancelado':
            s=next((y for y in rows if act(y) and y['dia']>=x['dia'] and abs(y['tot']-x['tot'])<0.01 and y['rfc']==x['rfc'] and y is not x),None)
            canc.append(dict(folio=x['folio'],dia=x['dia'],subtotal=x['sub'],iva=x['iva'],total=x['tot'],sustituto=(s['folio'] if s else None),sustituto_dia=(s['dia'] if s else None)))
    falt=[]; ser={}
    for x in rows:
        try: ser.setdefault(x['serie'],[]).append(int(x['folio']))
        except: pass
    for s,arr in ser.items():
        st=set(arr)
        if max(arr)-min(arr)<5000: falt+= [ (s+'-' if s else '')+str(n) for n in range(min(arr),max(arr)+1) if n not in st]
    X=r2(sum(c['base0'] for c in cand if c['fila'] in marc)); W=r2(sum(d['W_exacta'] for d in dias)); I=r2(sum(d['IEPS'] for d in dias))
    tot=dict(cfdi=len(rows),activos=sum(1 for x in rows if act(x)),cancelados=len(canc),dias=len(dias),
             S=r2(sum(d['S'] for d in dias)),T=r2(sum(d['T'] for d in dias)),U=r2(sum(d['U'] for d in dias)),V=r2(sum(d['V'] for d in dias)),W=W,X=X,IEPS=I,
             iva_explicado=r2((X-I)*0.16),no_explicada=r2(W-(X-I)*0.16),dias_revisar=sum(1 for d in dias if d['revision']=='REVISAR'),
             dias_con_W_mayor_tol=[d['dia'] for d in dias if abs(d['W_exacta'])>tol],base16=r2(sum(d['S'] for d in dias)-sum(d['V'] for d in dias)-X))
    for d in dias: d.pop('W_exacta')
    return dict(archivo=path.split('/')[-1],columna_resultados_inicia=max(last+1,19),ultima_columna_original=last,totales=tot,dias=dias,candidatos_sin_iva=cand,
                marcados_tasa0=sorted(c['folio'] for c in cand if c['fila'] in marc),ieps=ieps,cancelados=canc,folios_faltantes=falt)
if __name__=='__main__':
    # Cálculo independiente (sin el motor JS). Uso, desde la raíz del paquete:  python3 04_referencia_motor/golden.py
    import os
    root=os.path.join(os.path.dirname(os.path.abspath(__file__)),'..')
    for p,m in (('01_entradas/B_INGRESOS_2026-09_al_29.xlsx',('9856','9897')),('01_entradas/C_INGRESOS_2026-09_al_30.xlsx',())):
        g=golden(os.path.join(root,p),m); print(g['archivo'],json.dumps(g['totales'],ensure_ascii=False))
