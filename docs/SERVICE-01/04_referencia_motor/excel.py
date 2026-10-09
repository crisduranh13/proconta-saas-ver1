import pandas as pd, openpyxl, datetime as dt
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.formatting.rule import CellIsRule
R,O,A,m,rep,cf=pd.read_pickle('final.pkl')
F='Arial'; H=Font(name=F,bold=True,color='FFFFFF',size=10); HF=PatternFill('solid',fgColor='1F3864')
B=Font(name=F,size=10); BB=Font(name=F,size=10,bold=True); BLUE=Font(name=F,size=10,color='0000FF'); GREY=Font(name=F,size=9,italic=True,color='595959')
th=Side(style='thin',color='BFBFBF'); TOT=PatternFill('solid',fgColor='D9E1F2'); KEY=PatternFill('solid',fgColor='FFF2CC')
ACC='#,##0.00;[Red]-#,##0.00;-'
COL={'CONCILIADO':'C6EFCE','CONCILIADO C/OBS':'FFEB9C','PARCIAL':'FFD8A8','SIN CFDI':'FFC7CE','NO REQUIERE':'E7E6E6'}
W=[11,11,46,54,13,17,20,15,30,11,12,8,7,6,70]
wb=openpyxl.Workbook()
def hoja(ws,df,titulo=None):
    ws.append(list(df.columns)+['Diferencia'])
    for c in ws[1]: c.font=H; c.fill=HF; c.alignment=Alignment(wrap_text=True,vertical='center',horizontal='center')
    for r in df.values.tolist(): ws.append(r)
    n=ws.max_row
    for i in range(2,n+1):
        ws.cell(i,16,f'=IF(K{i}="","",E{i}-K{i})').number_format=ACC
        for j in range(1,17):
            c=ws.cell(i,j); c.font=B; c.border=Border(bottom=th); c.alignment=Alignment(wrap_text=True,vertical='top')
        ws.cell(i,5).number_format=ACC; ws.cell(i,11).number_format=ACC
        st=ws.cell(i,6).value
        if st in COL: ws.cell(i,6).fill=PatternFill('solid',fgColor=COL[st])
    ws.cell(n+1,4,'TOTAL').font=BB
    for j,col in ((5,'E'),(11,'K')):
        c=ws.cell(n+1,j,f'=SUM({col}2:{col}{n})'); c.number_format=ACC; c.font=BB; c.fill=TOT
    ws.cell(n+1,4).fill=TOT
    for i,w in enumerate(W+[12],1): ws.column_dimensions[L(i)].width=w
    ws.freeze_panes='A2'; ws.auto_filter.ref=f'A1:P{n}'
    ws.page_setup.orientation='landscape'; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
    ws.sheet_properties.pageSetUpPr.fitToPage=True; ws.print_title_rows='1:1'
    return n
ws=wb.active; ws.title='Conciliación salidas'; N=hoja(ws,R)
for e in ['CONCILIADO','CONCILIADO C/OBS','PARCIAL','SIN CFDI','NO REQUIERE']:
    nm={'CONCILIADO':'Conciliado','CONCILIADO C/OBS':'Conciliado con obs','PARCIAL':'Parcial','SIN CFDI':'Sin CFDI','NO REQUIERE':'No requiere CFDI'}[e]
    w=wb.create_sheet(nm); sub=R[R.Estatus==e].drop(columns=['Estatus'])
    w.append(list(sub.columns))
    for c in w[1]: c.font=H; c.fill=HF; c.alignment=Alignment(wrap_text=True,vertical='center',horizontal='center')
    for r in sub.values.tolist(): w.append(r)
    n=w.max_row
    for i in range(2,n+1):
        for j in range(1,15):
            c=w.cell(i,j); c.font=B; c.border=Border(bottom=th); c.alignment=Alignment(wrap_text=True,vertical='top')
        w.cell(i,5).number_format=ACC; w.cell(i,10).number_format=ACC
    w.cell(n+1,4,'TOTAL').font=BB; w.cell(n+1,4).fill=TOT
    for j,col in ((5,'E'),(10,'J')):
        c=w.cell(n+1,j,f'=SUM({col}2:{col}{n})'); c.number_format=ACC; c.font=BB; c.fill=TOT
    for i,wd in enumerate([11,11,46,54,13,20,15,30,11,12,8,7,6,70],1): w.column_dimensions[L(i)].width=wd
    w.freeze_panes='A2'; w.auto_filter.ref=f'A1:N{n}'; w.sheet_properties.tabColor=COL[e]
    w.page_setup.orientation='landscape'; w.page_setup.fitToWidth=1; w.page_setup.fitToHeight=0; w.sheet_properties.pageSetUpPr.fitToPage=True
# ---- CFDI sin pago identificado
usados=set(m.values())
nop=A[~A.cid.isin(usados)].sort_values('Fecha')
w=wb.create_sheet('CFDI sin pago')
w.append(['Fecha','RFC emisor','Emisor','Serie-Folio','Subtotal','IVA','Total','Método','Forma','Uso','Diagnóstico'])
for c in w[1]: c.font=H; c.fill=HF; c.alignment=Alignment(wrap_text=True,horizontal='center')
def diag(y):
    r=y['R.F.C. Emisor']
    if r=='GFE9707075U3': return 'Consumo del monedero de combustible: el abono al monedero ya salió del banco en otra fecha'
    if r=='OCM0505317M0': return 'Peaje: el CFDI se emite al cierre del mes por los consumos; los abonos salieron durante el mes'
    if r=='CGA010307N18' or r=='GSU010620SB9': return 'Combustible: revisar si se pagó con monedero, en efectivo o desde otra cuenta'
    if y['Método de Pago']=='PPD': return 'CFDI PPD: se acredita con el complemento de pago (REP). Verificar que exista el REP'
    if str(y['Forma de Pago'])=='1': return 'CFDI en efectivo: no hay salida bancaria. Si supera $2,000 no es deducible (art. 27 fr. III LISR)'
    if str(y['Forma de Pago'])=='4': return 'CFDI pagado con tarjeta de crédito: falta el estado de cuenta de la TDC'
    if str(y['Forma de Pago'])=='17': return 'CFDI con forma 17 (compensación): revisar contra cuentas por pagar'
    if str(y['Forma de Pago'])=='30': return 'CFDI con forma 30 (aplicación de anticipo): revisar el anticipo relacionado'
    return 'Sin salida bancaria identificada: confirmar forma de pago o buscar en otra cuenta'
for _,y in nop.iterrows():
    w.append([y.Fecha.strftime('%d/%m/%Y'),y['R.F.C. Emisor'],str(y['Razón Social Emisor'])[:40],y.SF,float(y['Sub Total']),float(y['Total IVA Tras.']),float(y.Total),y['Método de Pago'],'%02d'%int(float(y['Forma de Pago'])),y['Uso Cfdi Receptor'],diag(y)])
n=w.max_row
for i in range(2,n+1):
    for j in range(1,12):
        c=w.cell(i,j); c.font=B; c.border=Border(bottom=th); c.alignment=Alignment(wrap_text=True,vertical='top')
    for j in (5,6,7): w.cell(i,j).number_format=ACC
w.cell(n+1,4,'TOTAL').font=BB
for j,col in ((5,'E'),(6,'F'),(7,'G')):
    c=w.cell(n+1,j,f'=SUM({col}2:{col}{n})'); c.number_format=ACC; c.font=BB; c.fill=TOT
w.cell(n+1,4).fill=TOT
for i,wd in enumerate([11,15,40,20,13,12,13,8,7,6,78],1): w.column_dimensions[L(i)].width=wd
w.freeze_panes='A2'; w.auto_filter.ref=f'A1:K{n}'
NCFDI=n
pd.to_pickle(N,'N.pkl')
wb.save('base.xlsx'); print('ok',N,NCFDI,len(nop))
